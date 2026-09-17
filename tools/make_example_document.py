# -*- coding: utf-8 -*-
"""Создать обезличенный образец документа графика ТО/КМХ (xlsx).

Структура как у рабочего документа: лист «1» — график ТО («ТО-3 (20)»), лист «3» — дни
контроля (КМХ), лист «коды» — справочник с нормами времени и объединёнными ячейками.
Все данные вымышленные: коды, наименования, места установки и номера.

Запуск:
    .venv_311\Scripts\python.exe tools\make_example_document.py
    .venv_311\Scripts\python.exe tools\make_example_document.py --out uploads/example_tokmh.xlsx --rows 12
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from aila_core import config

MONTHS = ("январь", "февраль", "март", "апрель", "май", "июнь",
          "июль", "август", "сентябрь", "октябрь", "ноябрь", "декабрь")
TO_TYPES = ("ТО-1", "ТО-2", "ТО-3")
EQUIPMENT = (
    "Датчик давления, зав. № 000101",
    "Датчик температуры, зав. № 000102",
    "Преобразователь расхода, зав. № 000103",
    "Шкаф электроники, зав. № 000104",
    "Преобразователь плотности, зав. № 000105",
    "Блок обработки данных, зав. № 000106",
)
PLACES = ("Установка № 1", "Установка № 2", "Установка № 3")


def sheet_to(book: Any, rows: int) -> None:
    """Лист «1»: график ТО — в ячейке вид работ и день."""
    sheet = book.create_sheet("1")
    sheet["A1"], sheet["B1"], sheet["C1"], sheet["D1"] = (
        "№ п/п", "Код СИ/\nобору-\nдования", "Наименование СИ и оборудования", "Место установки")
    sheet["E1"] = "Вид (ТО-1, ТО-2, ТО-3) технического обслуживания"
    sheet.merge_cells(start_row=1, start_column=5, end_row=1, end_column=16)
    sheet.merge_cells("A1:A3")
    sheet.merge_cells("B1:B3")
    for index, month in enumerate(MONTHS):
        sheet.cell(row=2, column=5 + index, value=month)
        sheet.cell(row=3, column=5 + index, value="(число)")
    for offset in range(rows):
        row = 4 + offset
        sheet.cell(row=row, column=1, value=offset + 1)
        sheet.cell(row=row, column=2, value=offset % len(EQUIPMENT) + 1)
        sheet.cell(row=row, column=3, value=EQUIPMENT[offset % len(EQUIPMENT)])
        sheet.cell(row=row, column=4, value=PLACES[offset % len(PLACES)])
        for index in range(12):
            kind = TO_TYPES[(offset + index) % len(TO_TYPES)]
            day = (offset * 7 + index * 4) % 27 + 2
            sheet.cell(row=row, column=5 + index, value=f"{kind} ({day:02d})")


def sheet_kmx(book: Any, groups: int) -> None:
    """Лист «3»: дни контроля; код и наименование объединены по вертикали."""
    sheet = book.create_sheet("3")
    sheet["A1"], sheet["B1"] = "№ п/п", "Код\nСИ\n/\nобору\nдования"
    sheet["C1"], sheet["D1"] = "Место установки", "Наименование СИ и оборудования"
    sheet["F1"] = "Дата проведения контроля"
    sheet.merge_cells("A1:A3")
    sheet.merge_cells("B1:B3")
    sheet.merge_cells(start_row=1, start_column=6, end_row=1, end_column=17)
    for index, month in enumerate(MONTHS):
        sheet.cell(row=2, column=6 + index, value=month)
        sheet.cell(row=3, column=6 + index, value="(число)")
    row = 4
    number = 1
    for group in range(groups):
        size = 1 + group % 2                      # 1 или 2 строки на позицию
        for sub in range(size):                   # сначала значения, потом объединение
            sheet.cell(row=row + sub, column=1, value=number)
            sheet.cell(row=row + sub, column=2, value=group % len(EQUIPMENT) + 1)
            sheet.cell(row=row + sub, column=3, value=PLACES[group % len(PLACES)])
            sheet.cell(row=row + sub, column=4, value=f"Прибор контроля, зав. № {100 + group}")
            for index in range(12):
                days = sorted({(group * 7 + index * 3 + sub * 2) % 28 + 1,
                               (group * 5 + index * 4 + sub + 9) % 28 + 1})
                sheet.cell(row=row + sub, column=6 + index,
                           value=", ".join(f"{day:02d}" for day in days))
        for column in (1, 2, 3, 4):               # как в рабочем файле: объединённые ячейки
            sheet.merge_cells(start_row=row, start_column=column,
                              end_row=row + size - 1, end_column=column)
        row += size
        number += 1


def sheet_codes(book: Any, codes: int) -> None:
    """Лист «коды»: справочник с нормами времени, код и наименование объединены."""
    sheet = book.create_sheet("коды")
    headers = ("Код СИ/ оборудования)", "Наименование оборудования", "Вид ТО", "Кол-во ТО",
               "Нормы времени на одно ТО ВСЕГО", "", "Годовые нормы времени по видам ТО")
    for column, title in enumerate(headers, start=1):
        sheet.cell(row=1, column=column, value=title)
    sheet.merge_cells("C1:C2")
    sheet.merge_cells("E1:F2")
    sheet["D2"] = "в год"
    row = 3
    for code in range(1, codes + 1):
        start = row
        total = 0.0
        for kind, count, norm in (("ТО-1", 11, 1.0), ("ТО-2", 3, 6.0), ("ТО-3", 1, 20.0)):
            yearly = count * norm
            total += yearly
            sheet.cell(row=row, column=1, value=code)
            sheet.cell(row=row, column=2, value=EQUIPMENT[(code - 1) % len(EQUIPMENT)])
            sheet.cell(row=row, column=3, value=kind)
            sheet.cell(row=row, column=4, value=count)
            sheet.cell(row=row, column=5, value=norm)
            sheet.cell(row=row, column=7, value=yearly)
            row += 1
        sheet.cell(row=row, column=1, value=code)
        sheet.cell(row=row, column=2, value=EQUIPMENT[(code - 1) % len(EQUIPMENT)])
        sheet.cell(row=row, column=3, value="Всего в год")
        sheet.cell(row=row, column=7, value=total)
        row += 1
        for column in (1, 2):                     # объединяем код и наименование по группе
            sheet.merge_cells(start_row=start, start_column=column, end_row=row - 1, end_column=column)


def main() -> int:
    parser = argparse.ArgumentParser(description="Обезличенный образец документа ТО/КМХ")
    parser.add_argument("--out", default="uploads/example_tokmh.xlsx")
    parser.add_argument("--rows", type=int, default=12, help="строк на листе «1»")
    parser.add_argument("--groups", type=int, default=6, help="позиций на листе «3»")
    args = parser.parse_args()

    try:
        import openpyxl
    except ImportError:
        print("нужен openpyxl: .venv_311\\Scripts\\python.exe -m pip install openpyxl")
        return 2

    book = openpyxl.Workbook()
    book.remove(book.active)
    sheet_to(book, args.rows)
    sheet_kmx(book, args.groups)
    sheet_codes(book, 4)
    out = Path(args.out)
    if not out.is_absolute():
        out = config.ROOT / out
    out.parent.mkdir(parents=True, exist_ok=True)
    book.save(out)
    print("образец записан:", out, f"({out.stat().st_size / 1024:.1f} КБ)")
    print("листы: 1 (график ТО), 3 (дни контроля), коды (справочник с нормами)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
