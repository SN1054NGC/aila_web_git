# -*- coding: utf-8 -*-
"""RAG по документации: разбор Excel -> документы -> поиск.

По умолчанию используется BM25 на чистом Python:
  * НЕ требует моделей, интернета, torch и GPU;
  * мгновенно строится на сотнях строк и мгновенно ищет;
  * подходит для слабого ПК.
Опционально (AILA_RETRIEVAL=chroma) включается семантический поиск ChromaDB + e5-small,
но каталог с моделью должен лежать локально (models/e5-small).
"""
from __future__ import annotations

import json
import math
import re
from pathlib import Path
from typing import Any, Iterable

from . import compat  # noqa: F401  (важно: патч numpy до импорта chroma)
from . import config
from . import kmhto

# ---------------------------------------------------------------------------
# токенизация
# ---------------------------------------------------------------------------
_TOKEN_RE = re.compile(r"[0-9]+|[a-zа-яё]+", re.IGNORECASE)

_STOP = {
    "и", "в", "во", "не", "что", "он", "на", "я", "с", "со", "как", "а", "все", "она", "так",
    "его", "но", "да", "ты", "к", "у", "же", "вы", "за", "бы", "по", "только", "ее", "мне", "было",
    "вот", "от", "меня", "еще", "нет", "о", "из", "ему", "теперь", "когда", "даже", "ну", "ли",
    "если", "уже", "или", "ни", "быть", "был", "него", "до", "вас", "опять", "уж", "вам", "ведь",
    "там", "потом", "себя", "ничего", "ей", "может", "они", "тут", "где", "есть", "ней", "для",
    "мы", "тебя", "их", "чем", "была", "сам", "чтоб", "без", "будто", "чего", "раз", "тоже", "себе",
    "под", "будет", "тогда", "кто", "этот", "того", "потому", "этого", "какой", "совсем", "ним",
    "здесь", "этом", "один", "почти", "мой", "тем", "чтобы", "нее", "сейчас", "были", "куда",
    "зачем", "всех", "никогда", "можно", "при", "наконец", "два", "об", "другой", "хоть", "после",
    "над", "больше", "тот", "через", "эти", "нас", "про", "всего", "них", "какая", "много", "разве",
    "три", "эту", "моя", "хорошо", "свою", "этой", "перед", "иногда", "лучше", "чуть", "том",
    "нельзя", "такой", "им", "более", "всегда", "конечно", "всю", "между", "какое", "какие", "это",
    "эта", "этих", "скажи", "расскажи", "покажи", "нужно", "надо", "какой-нибудь", "какие-то",
    "the", "a", "an", "is", "are", "was", "were", "of", "to", "in", "and", "or", "on", "for", "at",
    "by", "with", "what", "which", "how", "when", "where", "who", "please", "tell", "show",
}

_RU_ENDINGS = tuple(sorted({
    "иями", "ями", "ами", "ией", "иях", "иям", "ием", "ого", "его", "ому", "ему", "ыми", "ими",
    "ая", "яя", "ое", "ее", "ые", "ие", "ой", "ей", "ый", "ий", "ом", "ем", "ам", "ям", "ах", "ях",
    "ов", "ев", "ую", "юю", "ия", "ию", "ии", "ье", "ья", "ью", "а", "я", "о", "е", "у", "ю", "ы",
    "и", "ь", "й",
}, key=len, reverse=True))


def _stem(token: str) -> str:
    w = token.lower().replace("ё", "е")
    if not w.isalpha():
        return w
    for end in _RU_ENDINGS:
        if len(w) - len(end) >= 4 and w.endswith(end):
            return w[: -len(end)]
    return w


_TO_TOKEN_RE = re.compile(r"ТО\s*[-–]?\s*([1-3])", re.IGNORECASE)

# Меняйте при правках разбора/токенизации: старый кэш станет недействительным.
INDEX_VERSION = 15


def tokenize(text: str) -> list[str]:
    text = text or ""
    out: list[str] = []
    for raw in _TOKEN_RE.findall(text):
        low = raw.lower().replace("ё", "е")
        if low in _STOP:
            continue
        out.append(low)
        if low.isalpha():
            stemmed = _stem(low)
            if stemmed != low:
                out.append(stemmed)
    # "ТО-1/2/3" — ключевой термин регламента. Общий токенизатор режет его на "то" и "3",
    # поэтому добавляем отдельный редкий токен "то3": он ищется точно и имеет высокий IDF.
    for num in _TO_TOKEN_RE.findall(text):
        out.append("то" + num)
    return out


def query_hints(query: str) -> dict[str, list[str]]:
    """Что именно ищет пользователь: код оборудования, тип ТО, номер строки."""
    return {
        "codes": re.findall(r"(?:код\w*|№|номер)\s*([0-9]{1,4})", query, re.IGNORECASE),
        "to": re.findall(r"ТО\s*[-–]?\s*([1-3])", query, re.IGNORECASE),
        "rows": re.findall(r"строк\w*\s*([0-9]{1,5})", query, re.IGNORECASE),
    }


class BM25:
    """Классический BM25 (k1, b) на чистом Python."""

    def __init__(self, docs_tokens: list[list[str]], k1: float = 1.5, b: float = 0.75) -> None:
        self.k1, self.b = k1, b
        self.n = len(docs_tokens)
        self.doc_len = [len(t) for t in docs_tokens]
        self.avgdl = (sum(self.doc_len) / self.n) if self.n else 0.0
        self.tf: list[dict[str, int]] = []
        df: dict[str, int] = {}
        for tokens in docs_tokens:
            counts: dict[str, int] = {}
            for t in tokens:
                counts[t] = counts.get(t, 0) + 1
            self.tf.append(counts)
            for t in counts:
                df[t] = df.get(t, 0) + 1
        self.idf = {t: math.log(1 + (self.n - d + 0.5) / (d + 0.5)) for t, d in df.items()}

    def scores(self, query_tokens: Iterable[str]) -> list[float]:
        result = [0.0] * self.n
        if not self.n or not self.avgdl:
            return result
        for token in set(query_tokens):
            idf = self.idf.get(token)
            if idf is None:
                continue
            for i, counts in enumerate(self.tf):
                f = counts.get(token)
                if not f:
                    continue
                norm = 1 - self.b + self.b * (self.doc_len[i] / self.avgdl)
                result[i] += idf * (f * (self.k1 + 1)) / (f + self.k1 * norm)
        return result


# ---------------------------------------------------------------------------
# разбор Excel
# ---------------------------------------------------------------------------
_TO_RE = re.compile(r"(ТО\s*[-–]?\s*[1-3])", re.IGNORECASE)
_CODE_COL_RE = re.compile(r"^код\b|\bкод\s+си|^код\s", re.IGNORECASE)
_NAME_COL_RE = re.compile(r"наимен|назван|^оборудован|средств", re.IGNORECASE)
_REF_SHEET_RE = re.compile(r"^код", re.IGNORECASE)


def _clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        if value.is_integer():
            return str(int(value))
    text = re.sub(r"\s+", " ", str(value)).strip()
    return "" if text.lower() in ("nan", "nat", "none", "null", "#n/a") else text


def _is_month_header(head: str) -> bool:
    return kmhto.is_month_header(head)


def _month_index(head: str) -> int:
    return kmhto.month_index(head)


def _is_units_row(row: list) -> bool:
    """Строка-пояснение шапки: "(число)", "в год" — это ещё не данные."""
    filled = [v for v in (_clean(x) for x in row) if v]
    if len(filled) < 2:
        return False
    for value in filled:
        if len(value) > 24:
            return False
        if re.search(r"\d", value) and not (value.startswith("(") and value.endswith(")")):
            return False
    return True


def _analyze_sheet(rows: list) -> tuple[list[str], int]:
    """Разбирает шапку листа и возвращает (заголовки колонок, номер первой строки данных).

    В регламентах СИКН шапка многоэтажная: группа -> названия месяцев -> "(число)".
    Название месяца важнее группы, потому что именно оно привязывает ТО к календарю.
    """
    width = max((len(r) for r in rows), default=0)
    head_row = 0
    while head_row < len(rows) and sum(1 for v in rows[head_row] if _clean(v)) < 2:
        head_row += 1
    if head_row >= len(rows):
        return [], len(rows)

    month_row, month_score = None, 0
    for i in range(head_row, min(len(rows), head_row + 6)):
        score = sum(1 for v in rows[i] if _is_month_header(_clean(v)))
        if score > month_score:
            month_row, month_score = i, score

    if month_row is None or month_score < 2:      # обычная односрочная шапка
        labels = []
        for col in range(len(rows[head_row])):
            value = _clean(rows[head_row][col])
            labels.append(value or f"Колонка{col + 1}")
        return labels, head_row + 1

    header_end = month_row
    if month_row + 1 < len(rows) and _is_units_row(rows[month_row + 1]):
        header_end = month_row + 1

    labels = []
    for col in range(width):
        collected: list[str] = []
        for row_idx in range(head_row, header_end + 1):
            if col < len(rows[row_idx]):
                value = _clean(rows[row_idx][col])
                if value:
                    collected.append(value)
        month = next((v for v in collected if _is_month_header(v)), None)
        label = month or (collected[-1] if collected else "")
        label = re.sub(r"\s+", " ", label).strip(" .-–")
        labels.append(label or f"Колонка{col + 1}")
    return labels, header_end + 1


def _detect_columns(headers: list[str]) -> tuple[int | None, int | None]:
    """Колонка кода и колонка наименования — строго по заголовкам."""
    code_col = name_col = None
    for idx, head in enumerate(headers):
        if not head:
            continue
        if code_col is None and _CODE_COL_RE.search(head):
            code_col = idx
            continue
        if name_col is None and _NAME_COL_RE.search(head):
            name_col = idx
    return code_col, name_col


# Опечатки и карта листов, найденные при последнем разборе (правила из проекта kmhto)
LAST_TYPOS: list[str] = []
LAST_SHEETS: list[dict[str, Any]] = []


def load_reference_docs(folder: Path | None = None) -> list[dict]:
    """Справочные материалы (data/reference/*.md) — знания о самом документе и предметной области.

    Каждый раздел «## Заголовок» становится отдельным документом, чтобы поиск находил
    конкретную тему, а не весь справочник целиком.
    """
    folder = folder or (config.DATA_DIR / "reference")
    if not folder.exists():
        return []
    docs: list[dict] = []
    for path in sorted(folder.glob("*.md")):
        try:
            text = path.read_text(encoding="utf-8")
        except OSError:
            continue
        title = path.stem
        current_title: str = title
        current_lines: list[str] = []
        for line in text.splitlines():
            if line.startswith("## "):
                if current_lines:
                    docs.append(_reference_doc(title, current_title, current_lines))
                current_title, current_lines = line[3:].strip(), []
            else:
                current_lines.append(line)
        if current_lines:
            docs.append(_reference_doc(title, current_title, current_lines))
    return [doc for doc in docs if len(doc["text"]) > 120]


def _reference_doc(source: str, title: str, lines: list[str]) -> dict:
    body = "\n".join(lines).strip()
    return {
        "id": f"ref:{source}:{title[:40]}",
        "kind": "reference",
        "text": f"СПРАВКА: {title}\nИсточник: {source}\n{body}",
        "meta": {"source": source, "sheet": "справка", "row": 0, "code": "",
                 "equipment": title, "months": {}, "ref": True},
    }


def read_merged_values(path: Path) -> dict[str, dict[tuple[int, int], Any]]:
    """Значения объединённых ячеек: лист -> {(строка, колонка): значение верхней ячейки}.

    Excel отдаёт значение только в левой верхней ячейке объединения, остальные пустые.
    Без этого строки «ТО-2», «ТО-3», «Всего в год» теряют код и оборудование.
    """
    import openpyxl  # тяжёлый импорт только здесь

    result: dict[str, dict[tuple[int, int], Any]] = {}
    book = openpyxl.load_workbook(path, data_only=True)
    for sheet in book.worksheets:
        filled: dict[tuple[int, int], Any] = {}
        for rng in sheet.merged_cells.ranges:
            value = sheet.cell(row=rng.min_row, column=rng.min_col).value
            if value is None:
                continue
            # Заполняем только ВНИЗ по первому столбцу объединения: так код и наименование
            # наследуются строками «ТО-2/ТО-3/Всего в год». Горизонтальные объединения
            # (одно значение на две колонки) не дублируем — иначе в тексте две копии.
            for row in range(rng.min_row + 1, rng.max_row + 1):
                filled[(row, rng.min_col)] = value
        result[str(sheet.title)] = filled
    return result


def fill_merged(rows: list[list[Any]], merges: dict[tuple[int, int], Any], data_start: int) -> None:
    """Разнести значения объединённых ячеек по строкам данных (шапку не трогаем)."""
    for (row_number, column), value in merges.items():
        row_idx, col_idx = row_number - 1, column - 1
        if row_idx < data_start:
            continue
        if 0 <= row_idx < len(rows) and 0 <= col_idx < len(rows[row_idx]):
            if rows[row_idx][col_idx] is None:
                rows[row_idx][col_idx] = value


def parse_workbook(path: Path, source_name: str | None = None) -> list[dict]:
    """Excel -> документы (строка = документ + сводка по каждому коду оборудования)."""
    import pandas as pd  # тяжёлый импорт только здесь

    path = Path(path)
    source = source_name or path.name
    sheets: dict[str, Any] = pd.read_excel(path, sheet_name=None, header=None)
    try:
        merged = read_merged_values(path)
    except Exception:            # openpyxl недоступен — сработает запасное наследование
        merged = {}
    docs: list[dict] = []
    by_code: dict[tuple[str, str], dict[str, Any]] = {}
    # нормы времени из листа-справочника: (лист, код) -> строки с нормами
    ref_norms: dict[tuple[str, str], list[tuple[int, str, list[tuple[str, str]]]]] = {}
    all_typos: list[str] = []
    sheets_info: list[dict[str, Any]] = []
    LAST_TYPOS.clear()

    for sheet_name, df in sheets.items():
        if df is None or df.empty:
            continue
        # pandas через where(..., None) оставляет NaN в числовых колонках — приводим к None,
        # иначе заполнение объединённых ячеек молча не срабатывает.
        rows = [[None if value is None or (isinstance(value, float) and math.isnan(value)) else value
                 for value in row] for row in df.values.tolist()]
        if not rows:
            continue
        headers, data_start = _analyze_sheet(rows)
        # Свои (не унаследованные) значения строки: пустые строки внутри объединённых блоков
        # не должны становиться отдельными документами.
        own_values = [sum(1 for value in row if _clean(value)) for row in rows]
        # Объединённые ячейки заполняем ТОЛЬКО в данных: шапка уже разобрана выше,
        # иначе её значения «поедут» в строки-пояснения.
        fill_merged(rows, merged.get(str(sheet_name), {}), data_start)
        code_col, name_col = _detect_columns(headers)
        month_cols = {col: _month_index(head) for col, head in enumerate(headers) if _month_index(head)}
        # Роль листа (1 — график ТО, 3 — график КМХ, data/коды — справочник)
        role = kmhto.sheet_role(str(sheet_name), headers)
        sheet_typos: list[str] = []

        sheet_rows = sheet_codes = 0
        fill_code_raw = fill_equipment = ""
        fill_ok = False

        for offset, row in enumerate(rows[data_start:], start=data_start):
            cells = [_clean(v) for v in row]
            if sum(1 for c in cells if c) < 2:
                continue
            if own_values[offset] == 0:
                continue                 # строка целиком из объединённых ячеек — не данные
            excel_row = offset + 1
            sheet_rows += 1

            pairs: list[tuple[str, str]] = []
            for col, value in enumerate(cells):
                if not value or col in month_cols:
                    continue                 # месяцы уходят отдельной строкой ниже
                if col in (code_col, name_col):
                    continue                 # код и наименование уже выведены выше
                head = headers[col] if col < len(headers) else f"Колонка{col + 1}"
                if _CODE_COL_RE.search(head) or _NAME_COL_RE.search(head) or re.match(r"^№", head):
                    continue                 # служебные колонки шапки
                pairs.append((head, value[:240]))

            code_raw = cells[code_col][:40] if code_col is not None and code_col < len(cells) else ""
            equipment = cells[name_col][:200] if name_col is not None and name_col < len(cells) else ""
            # В файле ячейки кода и наименования объединены по вертикали: строки «ТО-2»,
            # «ТО-3», «Всего в год» идут без кода. Наследуем их от строки-владельца, пока
            # не начался новый блок без кода. Иначе эти строки теряют привязку к оборудованию.
            if not code_raw and not equipment:
                if fill_ok:
                    code_raw, equipment = fill_code_raw, fill_equipment
            elif code_raw:
                fill_code_raw, fill_equipment, fill_ok = code_raw, equipment, True
            else:
                fill_ok = False            # блок без кода — дальше не подставляем
            code = ""
            if code_raw:
                digits = re.sub(r"[^0-9]", "", code_raw)
                if digits and len(digits) <= 6:
                    code = digits          # индекс/подсказки работают по цифрам

            months: dict[str, str] = {}
            month_days: dict[str, list[int]] = {}
            month_works: dict[str, list[list[Any]]] = {}
            for col, value in enumerate(cells):
                if not value or col not in month_cols:
                    continue
                month = headers[col]
                months[month] = value
                # Разбор по правилам проекта kmhto: «грязные» ячейки нормализуются,
                # недопустимые дни (0 и > 31) отбрасываются.
                days = kmhto.parse_days(value)
                if days:
                    month_days[month] = days
                    if re.search(r"\b(0|3[2-9]|[4-9]\d)\b", value):
                        sheet_typos.append(f"{sheet_name} строка {excel_row}, {month}: "
                                           f"недопустимый день — {value[:40]}")
                works = [[kind, day] for kind, day in kmhto.parse_work_items(value)] if role == "to" \
                    else []
                if works:
                    month_works[month] = works
                elif role == "to":
                    sheet_typos.append(f"{sheet_name} строка {excel_row}, {month}: "
                                       f"не разобран вид ТО — {value[:40]}")

            if month_works:
                schedule = "; ".join(
                    f"{month}: " + ", ".join(
                        f"{kind} ({day})" if day else kind for kind, day in works)
                    for month, works in month_works.items())
                caption = "Работы по месяцам (вид ТО и день)"
            elif month_days:
                schedule = "; ".join(f"{month}: " + ", ".join(str(d) for d in days)
                                     for month, days in month_days.items())
                caption = "Дни контроля по месяцам"
            else:
                schedule = ""
                caption = ""

            lines = [f"Лист: {sheet_name}", f"Строка Excel: {excel_row}"]
            if code_raw:
                lines.append(f"Код оборудования: {code_raw}")
            if equipment:
                lines.append(f"Оборудование: {equipment}")
            lines.extend(f"{head}: {value}" for head, value in pairs)
            if schedule:
                lines.append(f"{caption}: {schedule}")

            docs.append({
                "id": f"{sheet_name}:{excel_row}",
                "kind": "row",
                "text": "\n".join(lines),
                "meta": {
                    "source": source,
                    "sheet": str(sheet_name),
                    "row": excel_row,
                    "code": code,
                    "code_raw": code_raw,
                    "equipment": equipment,
                    "months": months,
                    "days": month_days,
                    "works": month_works,
                    "role": role,
                    "ref": bool(_REF_SHEET_RE.search(str(sheet_name))),
                },
            })

            if code:
                sheet_codes += 1
            if role == "ref" and code and pairs:
                ref_norms.setdefault((str(sheet_name), code), []).append((excel_row, equipment, pairs))
            if code and months:
                info = by_code.setdefault((str(sheet_name), code),
                                          {"equipment": equipment, "rows": [],
                                           "types": set(), "months": {}})
                info["rows"].append(excel_row)
                if not info["equipment"]:
                    info["equipment"] = equipment
                for value in months.values():
                    match = _TO_RE.search(value)
                    if match:
                        info["types"].add(re.sub(r"\s+", "", match.group(1)).upper())
                for month, value in months.items():
                    info["months"].setdefault(month, []).append(value)

        # Профиль листа и найденные опечатки — «информация по документу»
        all_typos.extend(sheet_typos)
        sheets_info.append({
            "name": str(sheet_name),
            "role": role,
            "rows": sheet_rows,
            "codes": sheet_codes,
            "months": len(month_cols),
            "month_start": kmhto.column_letter(min(month_cols)) if month_cols else "",
            "year": kmhto.document_year(rows, range(data_start)),
            "columns": [head for head in headers[:6] if head],
        })

    # Опечатки и карта листов: это и есть «информация по документу»,
    # её получает модель и показывает интерфейс.
    LAST_TYPOS[:] = all_typos
    if sheets_info:
        LAST_SHEETS[:] = sheets_info

    # сводка по каждому коду оборудования — отвечает на вопросы «код 3»
    for (sheet_name, code), info in sorted(
            by_code.items(), key=lambda kv: (len(kv[0][1]), kv[0][1], kv[0][0])):
        parts: list[str] = []
        for month_name, values in info["months"].items():
            unique = sorted(set(values))
            if not unique:
                continue
            if info["types"]:
                # "январь: ТО-3 (07); январь: ТО-3 (20)" — модель не путает месяцы
                parts.extend(f"{month_name}: {value}" for value in unique)
            else:
                parts.append(f"{month_name}: {', '.join(unique)}")
        months_text = "; ".join(parts)
        if info["types"]:
            caption = "График ТО по месяцам"
        else:
            caption = "Даты контроля по месяцам"
        text = (f"Лист: {sheet_name}\nКод оборудования: {code}\n"
                f"Оборудование: {info['equipment']}\n"
                f"Всего позиций с этим кодом: {len(info['rows'])}\n"
                f"Типы ТО по коду: {', '.join(sorted(info['types'])) or 'не указаны'}\n"
                f"{caption}: {months_text}")
        docs.append({
            "id": f"code:{sheet_name}:{code}",
            "kind": "code_summary",
            "text": text,
            "meta": {"source": source, "sheet": str(sheet_name), "row": 0, "code": code,
                     "equipment": info["equipment"], "months": info["months"],
                     "ref": bool(_REF_SHEET_RE.search(str(sheet_name))), "role": role},
        })

    # Сводка норм времени по коду (лист-справочник): отвечает на «какая норма времени
    # на ТО-3 у кода 3» и связывает строки с объединёнными ячейками с оборудованием.
    for (sheet_name, code), items in sorted(ref_norms.items(), key=lambda kv: (len(kv[0][1]), kv[0][1])):
        name = next((equipment for _row, equipment, _pairs in items if equipment), "")
        norms = next((pairs for _row, _equipment, pairs in items if pairs), [])
        lines = [f"Лист: {sheet_name}", f"Код оборудования: {code}"]
        if name:
            lines.append(f"Оборудование: {name}")
        lines.append(f"Строк с нормами времени: {len(items)}")
        if norms:
            lines.append("Нормы времени по видам работ: " + ", ".join(
                f"{head}: {value}" for head, value in norms))
        lines.append("Все нормы по коду:")
        for row_number, _equipment, pairs in items:
            lines.append(f"  строка {row_number}: " + ", ".join(f"{head}: {value}" for head, value in pairs))
        docs.append({
            "id": f"ref:{sheet_name}:{code}",
            "kind": "code_ref",
            "text": "\n".join(lines),
            "meta": {"source": source, "sheet": str(sheet_name), "row": 0, "code": code,
                     "equipment": name, "months": {}, "ref": False, "role": "ref"},
        })
    return docs


# ---------------------------------------------------------------------------
# движок
# ---------------------------------------------------------------------------
class RagEngine:
    """Поиск по проиндексированным документам (BM25; опционально ChromaDB)."""

    def __init__(self) -> None:
        self.docs: list[dict] = []
        self.bm25: BM25 | None = None
        self.source: str = ""
        self.source_path: str = ""
        # Типы из сторонних библиотек — аннотируем Any, чтобы не тянуть их в статический анализ
        self.chroma_collection: Any = None
        self.chroma_client: Any = None
        self.embedder: Any = None
        self.mode: str = "bm25"
        self.error: str = ""
        self.typos: list[str] = []
        self.sheets_info: list[dict[str, Any]] = []

    # ---------------------------------------------------------------- индекс
    def build(self, path: Path, source_name: str | None = None) -> dict:
        path = Path(path)
        if not path.exists():
            raise FileNotFoundError(f"Файл не найден: {path}")
        docs = parse_workbook(path, source_name or path.name)
        if not docs:
            raise ValueError("В файле не найдено ни одной строки с данными")
        docs = docs + load_reference_docs()      # знания о документе и предметной области
        self.docs = docs
        self.bm25 = BM25([tokenize(d["text"]) for d in docs])
        self.source = source_name or path.name
        self.source_path = str(path)
        self.error = ""
        self.mode = "bm25"
        self.typos = list(LAST_TYPOS)
        self.sheets_info = list(LAST_SHEETS)
        self._save_cache()
        if config.RETRIEVAL == "chroma":
            try:
                self._build_chroma()
            except Exception as exc:
                self.error = f"Chroma недоступна, используется BM25: {exc}"
        return self.status()

    def _save_cache(self) -> None:
        try:
            config.INDEX_CACHE.write_text(json.dumps(
                {"version": INDEX_VERSION, "source": self.source,
                 "source_path": self.source_path, "docs": self.docs,
                 "typos": self.typos, "sheets_info": self.sheets_info},
                ensure_ascii=False), encoding="utf-8")
        except OSError:
            pass

    def load_cache(self) -> bool:
        """Быстрый старт: индекс поднимается с диска, без повторного парсинга Excel."""
        try:
            if not config.INDEX_CACHE.exists():
                return False
            payload = json.loads(config.INDEX_CACHE.read_text(encoding="utf-8"))
            if payload.get("version") != INDEX_VERSION:
                return False        # кэш от старой версии разбора — перестроим
            docs = payload.get("docs") or []
            if not docs:
                return False
            self.docs = docs
            self.bm25 = BM25([tokenize(d["text"]) for d in docs])
            self.source = payload.get("source", "")
            self.source_path = payload.get("source_path", "")
            self.typos = list(payload.get("typos") or [])
            self.sheets_info = list(payload.get("sheets_info") or [])
            return True
        except Exception as exc:
            self.error = f"не удалось прочитать кэш индекса: {exc}"
            return False

    def ensure_loaded(self, default_xlsx: Path | None = None) -> None:
        if self.docs:
            return
        if not self.load_cache() and default_xlsx and Path(default_xlsx).exists():
            try:
                self.build(Path(default_xlsx))
            except Exception as exc:
                self.error = str(exc)

    # ---------------------------------------------------------------- chroma
    def _build_chroma(self) -> None:
        import chromadb
        from chromadb.config import Settings

        if not (config.EMBED_MODEL_DIR.exists() and any(config.EMBED_MODEL_DIR.iterdir())):
            raise RuntimeError(
                f"нет локальной модели эмбеддингов в {config.EMBED_MODEL_DIR} "
                "(см. tools/fetch_models.py)")
        from sentence_transformers import SentenceTransformer
        self.embedder = SentenceTransformer(str(config.EMBED_MODEL_DIR), device="cpu")
        self.chroma_client = chromadb.PersistentClient(
            path=str(config.RAG_DB_DIR), settings=Settings(anonymized_telemetry=False))
        try:
            # ВАЖНО: delete_collection вместо delete(where={}), который в chroma 0.4.x всегда падает
            self.chroma_client.delete_collection(config.CHROMA_COLLECTION)
        except Exception:
            pass
        self.chroma_collection = self.chroma_client.create_collection(
            config.CHROMA_COLLECTION, metadata={"hnsw:space": "cosine"})
        # e5 требует РАЗНЫЕ префиксы: passage: для документов, query: для запросов
        vectors = self.embedder.encode(
            [f"passage: {d['text']}" for d in self.docs],
            normalize_embeddings=True, batch_size=32, show_progress_bar=False)
        self.chroma_collection.add(
            ids=[d["id"] for d in self.docs],
            documents=[d["text"] for d in self.docs],
            embeddings=[v.tolist() for v in vectors],
            metadatas=[{k: (v if isinstance(v, (str, int, float)) else str(v))
                        for k, v in d["meta"].items()} for d in self.docs],
        )
        self.mode = "chroma"

    # ---------------------------------------------------------------- поиск
    def search(self, query: str, k: int | None = None) -> list[dict]:
        k = k or config.TOP_K
        if self.mode == "chroma" and self.chroma_collection is not None and self.embedder:
            return self._search_chroma(query, k)
        return self._search_bm25(query, k)

    def _search_bm25(self, query: str, k: int) -> list[dict]:
        if not self.bm25 or not self.docs:
            return []
        raw = self.bm25.scores(tokenize(query))
        best = max(raw) if raw else 0.0
        if best <= 0:
            return []
        hints = query_hints(query)
        order = sorted(range(len(raw)), key=lambda i: raw[i], reverse=True)

        ranked: list[tuple[float, dict]] = []
        for idx in order[: max(k * 8, 60)]:
            if raw[idx] <= 0:
                continue
            base = raw[idx] / best
            if base < 0.12:
                break                      # список отсортирован: дальше только шум
            doc = self.docs[idx]
            meta = doc["meta"]
            kind = doc["kind"]
            mult = 1.0
            if kind in ("code_summary", "code_ref"):
                mult *= 3.0 if meta.get("code") in hints["codes"] else 0.9
            if kind == "row" and meta.get("code") and meta["code"] in hints["codes"]:
                mult *= 1.8
            query_low = query.lower()
            # Справочные разделы важнее строк документа, когда спрашивают про понятия
            if kind == "reference" and not hints["codes"] and any(word in query_low for word in (
                    "что такое", "чем отлича", "расшифруй", "какие бывают", "объясни",
                    "для чего", "зачем", "термин", "норма времени", "технологическая карта",
                    "означа", "значит", "смысл", "сколько", "перечисли", "какие поля",
                    "структур", "формат", "описан", "правила", "как понять")):
                # Справочник описывает понятия словами, поэтому для определений он важнее строк
                mult *= 3.2
            asks_schedule = bool(hints["to"]) or any(
                word in query_low for word in ("график", "когда", "дата", "периодичн"))
            # «работы/обслуживание» — это график ТО, «контроль/дни/поверка» — график КМХ
            role = str(meta.get("role") or "")
            asks_kmx = any(word in query_low for word in
                           ("контрол", "кмх", "поверк", "калибр", "дни"))
            asks_to = any(word in query_low for word in
                          ("работ", "обслужив", "то-1", "то-2", "то-3", "то "))
            # «нормы времени» есть только в листе-справочнике: строки графиков тут не помогут
            if "норм" in query_low:
                if kind == "code_ref" or role == "ref":
                    mult *= 2.2
                else:
                    mult *= 0.5
            if role == "kmx" and asks_kmx and not asks_to:
                mult *= 1.7
            elif role == "to" and asks_to and not asks_kmx:
                mult *= 1.7
            if asks_schedule and kind == "row":
                months = meta.get("months") or {}
                if months and (not hints["to"] or any(
                        f"ТО-{n}" in str(v).replace(" ", "")
                        for n in hints["to"] for v in months.values())):
                    mult *= 1.7
            if hints["rows"] and str(meta.get("row")) in hints["rows"]:
                mult *= 1.6
            if meta.get("ref"):
                mult *= 0.55               # справочные листы менее полезны для ответа
            if kind == "row" and meta.get("equipment") and meta.get("months"):
                mult *= 1.1                # полноценная строка оборудования
            ranked.append((base * mult, doc))

        ranked.sort(key=lambda item: item[0], reverse=True)

        hits: list[dict] = []
        seen: set[str] = set()
        for score, doc in ranked:
            meta = doc["meta"]
            key = (f"{meta.get('equipment')}|{meta.get('code')}") if meta.get("equipment") \
                else f"{doc['id']}"
            if key in seen:
                continue
            seen.add(key)
            hits.append({
                "id": doc["id"],
                "text": doc["text"],
                "meta": meta,
                "kind": doc["kind"],
                "score": round(score, 4),
            })
            if len(hits) >= k:
                break

        if hits:
            top = hits[0]["score"] or 1.0
            for hit in hits:               # нормализуем 0..1 уже для показа
                hit["score"] = round(min(1.0, hit["score"] / top), 4)
        return hits

    def _search_chroma(self, query: str, k: int) -> list[dict]:
        vector = self.embedder.encode([f"query: {query}"], normalize_embeddings=True)[0]
        res = self.chroma_collection.query(query_embeddings=[vector.tolist()], n_results=k)
        hits = []
        for i, doc in enumerate(res.get("documents", [[]])[0]):
            dist = res.get("distances", [[0]])[0][i]
            hits.append({
                "id": res.get("ids", [[""]])[0][i],
                "text": doc,
                "meta": res.get("metadatas", [[{}]])[0][i] or {},
                "kind": "row",
                "score": round(max(0.0, 1.0 - float(dist)), 4),
            })
        return hits[:k]

    def build_context(self, hits: list[dict], max_chars: int = 3500) -> str:
        if not hits:
            return ""
        parts: list[str] = []
        used = 0
        for i, hit in enumerate(hits, start=1):
            meta = hit["meta"]
            where = []
            if meta.get("sheet"):
                where.append(f"лист {meta['sheet']}")
            if meta.get("row"):
                where.append(f"строка {meta['row']}")
            head = f"[{i}] " + (", ".join(where) if where else "сводка по коду")
            chunk = f"{head}\n{hit['text'][:900]}"
            if used + len(chunk) > max_chars:
                break
            parts.append(chunk)
            used += len(chunk)
        return "\n\n---\n\n".join(parts)

    # ---------------------------------------------------------------- статистика
    def digest(self) -> str:
        """Короткая сводка о документе: идёт в system-промпт для вопросов «сколько всего»."""
        if not self.docs:
            return ""
        rows = [d for d in self.docs if d["kind"] == "row"]
        codes = {d["meta"].get("code") for d in rows if d["meta"].get("code")}
        sheets = sorted({d["meta"].get("sheet", "") for d in rows if d["meta"].get("sheet")})
        by_sheet: dict[str, int] = {}
        for doc in rows:
            key = doc["meta"].get("sheet", "?")
            by_sheet[key] = by_sheet.get(key, 0) + 1
        role_names = {"to": "график ТО", "kmx": "график КМХ", "ref": "справочник", "other": "данные"}
        sheet_roles: dict[str, str] = {}
        for doc in rows:
            key = str(doc["meta"].get("sheet", "?"))
            sheet_roles.setdefault(key, str(doc["meta"].get("role") or "other"))
        detail = "; ".join(
            f"лист {key} ({role_names.get(sheet_roles.get(key, 'other'), 'данные')}): {count} строк"
            for key, count in sorted(by_sheet.items()))
        facts = (f"ФАКТЫ О ДОКУМЕНТЕ: файл {self.source or '—'}; "
                 f"всего строк данных — {len(rows)} ({detail}); "
                 f"уникальных кодов оборудования — {len(codes)}; "
                 f"листы — {', '.join(sheets) or '—'}. "
                 "Эти числа — истина, не считай позиции по фрагментам.")
        card = kmhto.describe_document(self.sheets_info)
        result = facts + "\n\n" + kmhto.DOCUMENT_INFO
        if card:
            result += "\n\n" + card
        if self.typos:
            result += ("\n\nВ документе есть непонятные ячейки (первые): "
                       + "; ".join(self.typos[:3]))
        return result

    # ---------------------------------------------------------------- статус
    def status(self) -> dict:
        sheets = sorted({d["meta"].get("sheet", "") for d in self.docs if d["meta"].get("sheet")})
        codes = {d["meta"].get("code") for d in self.docs if d["meta"].get("code")}
        return {
            "ready": bool(self.docs),
            "mode": self.mode,
            "docs": len(self.docs),
            "source": self.source,
            "sheets": sheets,
            "codes": len(codes),
            "typos": len(self.typos),
            "typo_examples": self.typos[:3],
            "sheets_info": self.sheets_info,
            "error": self.error,
        }

    def clear(self) -> dict:
        self.docs = []
        self.bm25 = None
        self.source = ""
        self.source_path = ""
        self.mode = "bm25"
        try:
            if config.INDEX_CACHE.exists():
                config.INDEX_CACHE.unlink()
        except OSError:
            pass
        try:
            if self.chroma_client is not None:
                self.chroma_client.delete_collection(config.CHROMA_COLLECTION)
                self.chroma_collection = None
        except Exception:
            pass
        return self.status()


if __name__ == "__main__":
    import sys

    eng = RagEngine()
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else config.sample_document()
    if target is None:
        raise SystemExit("укажите файл: python -m aila_core.rag <путь.xlsx>")
    print("build:", eng.build(target))
    for q in ["какое оборудование под кодом 3", "график ТО преобразователя расхода",
              "когда проводят ТО-3", "поверка влагомера", "шкаф ИВК СИКН"]:
        print(f"\n### {q}")
        for h in eng.search(q, 3):
            print("  ", h["score"], h["id"], "|",
                  " / ".join(h["text"].splitlines()[:4])[:120])
